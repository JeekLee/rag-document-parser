from __future__ import annotations

import io
import subprocess
import sys
import zipfile
from base64 import b64decode
from datetime import date

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill


PNG_BYTES = b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
    "YAAAAAYAAjCB0C8AAAAASUVORK5CYII="
)


def _workbook_bytes(workbook: Workbook) -> bytes:
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _rewrite_xlsx(
    data: bytes,
    rewrite,
) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source:
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():
                name, payload = rewrite(info.filename, source.read(info.filename))
                target.writestr(name, payload)
    return output.getvalue()


def _only_table(parsed):
    tables = [unit for unit in parsed.units if unit.type == "table"]
    assert len(tables) == 1
    return tables[0]


def test_xlsx_backend_separates_preamble_text_from_single_header_table():
    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Summary"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "Quarterly report"
    sheet["A2"] = "Period"
    sheet["B2"] = "2026 Q2"
    sheet.append([])
    sheet.append(["Item", "Amount", "Approved"])
    sheet.append(["Clinic", 1000, True])
    sheet.append(["Hospital", 2500, False])

    parsed = XlsxBackend().parse(_workbook_bytes(workbook), ".xlsx")

    assert XlsxBackend.supported_suffixes == (".xlsx",)
    assert [unit.type for unit in parsed.units] == ["text", "text", "table"]
    assert [unit.content for unit in parsed.units[:2]] == [
        "Quarterly report",
        "Period | 2026 Q2",
    ]
    assert all(
        unit.metadata["common"]["section_path"] == ["Summary"] for unit in parsed.units
    )

    table = parsed.units[2]
    assert table.format == "structured_table"
    assert [column["text"] for column in table.content["columns"]] == [
        "Item",
        "Amount",
        "Approved",
    ]
    assert table.metadata["spreadsheet"] == {
        "sheet_name": "Summary",
        "sheet_state": "visible",
        "cell_range": "A4:C6",
        "header_range": "A4:C4",
        "data_range": "A5:C6",
    }
    assert [row["source_row"] for row in table.content["rows"]] == [5, 6]
    assert [cell["address"] for cell in table.content["rows"][0]["cells"]] == [
        "A5",
        "B5",
        "C5",
    ]
    assert [cell["text"] for cell in table.content["rows"][0]["cells"]] == [
        "Clinic",
        "1000",
        "TRUE",
    ]
    assert table.content["rows"][1]["cells"][2]["text"] == "FALSE"


def test_xlsx_backend_flattens_merged_multirow_headers_and_preserves_spans():
    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.merge_cells("A1:A2")
    sheet["A1"] = "Region"
    sheet.merge_cells("B1:C1")
    sheet["B1"] = "Sales"
    sheet["B2"] = "January"
    sheet["C2"] = "February"
    sheet["A3"] = "Seoul"
    sheet["B3"] = 10
    sheet["C3"] = 20
    sheet["B4"] = 11
    sheet["C4"] = 21
    sheet.merge_cells("A3:A4")

    parsed = XlsxBackend().parse(_workbook_bytes(workbook), ".xlsx")

    table = _only_table(parsed)
    assert [column["text"] for column in table.content["columns"]] == [
        "Region",
        "Sales / January",
        "Sales / February",
    ]
    assert table.metadata["spreadsheet"] == {
        "sheet_name": "Sales",
        "sheet_state": "visible",
        "cell_range": "A1:C4",
        "header_range": "A1:C2",
        "data_range": "A3:C4",
    }

    header_rows = table.content["header_rows"]
    assert [row["source_row"] for row in header_rows] == [1, 2]
    assert [
        (cell["address"], cell["text"], cell["rowspan"], cell["colspan"])
        for cell in header_rows[0]["cells"]
    ] == [
        ("A1", "Region", 2, 1),
        ("B1", "Sales", 1, 2),
    ]
    assert [cell["address"] for cell in header_rows[1]["cells"]] == ["B2", "C2"]

    rows = table.content["rows"]
    assert [row["source_row"] for row in rows] == [3, 4]
    assert [
        (cell["address"], cell["text"], cell["rowspan"], cell["colspan"])
        for cell in rows[0]["cells"]
    ] == [
        ("A3", "Seoul", 2, 1),
        ("B3", "10", 1, 1),
        ("C3", "20", 1, 1),
    ]
    assert [cell["address"] for cell in rows[1]["cells"]] == ["B4", "C4"]


def test_xlsx_backend_falls_back_to_formula_text_when_cache_is_missing():
    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Formula"
    sheet.append(["Item", "Base", "Computed"])
    sheet.append(["Plan A", 5, "=B2*2"])

    parsed = XlsxBackend().parse(_workbook_bytes(workbook), ".xlsx")

    table = _only_table(parsed)
    formula_cell = table.content["rows"][0]["cells"][2]
    assert formula_cell["text"] == "=B2*2"
    assert formula_cell["address"] == "C2"
    assert formula_cell["formula"] == "=B2*2"
    assert "number_format" not in formula_cell
    assert "hyperlink" not in formula_cell

    assert len(parsed.quality_warnings) == 1
    warning = parsed.quality_warnings[0]
    warning_contract = {
        key: warning[key]
        for key in ("type", "severity", "sheet_name", "cell", "formula")
    }
    assert warning_contract == {
        "type": "xlsx_formula_cache_missing",
        "severity": "medium",
        "sheet_name": "Formula",
        "cell": "C2",
        "formula": "=B2*2",
    }


def test_xlsx_backend_prefers_cached_formula_value_when_available():
    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Formula"
    sheet.append(["Item", "Base", "Computed"])
    sheet.append(["Plan A", 5, "=B2*2"])

    def rewrite(name: str, payload: bytes) -> tuple[str, bytes]:
        if name == "xl/worksheets/sheet1.xml":
            payload = payload.replace(
                b"<f>B2*2</f><v />",
                b"<f>B2*2</f><v>10</v>",
            )
        return name, payload

    data = _rewrite_xlsx(_workbook_bytes(workbook), rewrite)
    parsed = XlsxBackend().parse(data, ".xlsx")

    formula_cell = _only_table(parsed).content["rows"][0]["cells"][2]
    assert formula_cell["text"] == "10"
    assert formula_cell["formula"] == "=B2*2"
    assert parsed.quality_warnings == []


def test_xlsx_backend_formats_typed_values_and_preserves_cell_semantics():
    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Typed values"
    sheet.append(["Date", "Active", "Rate", "Code", "Policy"])
    sheet.append([date(2026, 6, 4), True, 0.25, 7, "Guidance"])
    sheet["A2"].number_format = "yyyy-mm-dd"
    sheet["C2"].number_format = "0.00%"
    sheet["D2"].number_format = "0000"
    sheet["E2"].hyperlink = "https://example.test/policy"

    parsed = XlsxBackend().parse(_workbook_bytes(workbook), ".xlsx")

    table = _only_table(parsed)
    cells = table.content["rows"][0]["cells"]
    assert [cell["text"] for cell in cells] == [
        "2026-06-04",
        "TRUE",
        "25.00%",
        "0007",
        "Guidance",
    ]
    assert [cell["address"] for cell in cells] == ["A2", "B2", "C2", "D2", "E2"]
    assert cells[0]["number_format"] == "yyyy-mm-dd"
    assert "number_format" not in cells[1]
    assert cells[2]["number_format"] == "0.00%"
    assert cells[3]["number_format"] == "0000"
    assert cells[4]["hyperlink"] == "https://example.test/policy"
    assert "number_format" not in cells[4]
    assert all("formula" not in cell for cell in cells)


def test_xlsx_backend_ignores_style_only_cells_outside_the_value_range():
    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Trimmed"
    sheet.append(["Item", "Value"])
    sheet.append(["A", 1])
    sheet["Z100"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")

    parsed = XlsxBackend().parse(_workbook_bytes(workbook), ".xlsx")

    table = _only_table(parsed)
    assert table.metadata["spreadsheet"] == {
        "sheet_name": "Trimmed",
        "sheet_state": "visible",
        "cell_range": "A1:B2",
        "header_range": "A1:B1",
        "data_range": "A2:B2",
    }
    assert len(table.content["columns"]) == 2
    assert len(table.content["rows"]) == 1
    assert [cell["address"] for cell in table.content["rows"][0]["cells"]] == [
        "A2",
        "B2",
    ]


def test_xlsx_backend_keeps_hidden_sheets_with_sheet_state_metadata():
    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    workbook.active.title = "Visible"
    hidden = workbook.create_sheet("Archive")
    hidden.sheet_state = "hidden"
    hidden.append(["Item", "Value"])
    hidden.append(["Old", 9])

    parsed = XlsxBackend().parse(_workbook_bytes(workbook), ".xlsx")

    table = _only_table(parsed)
    assert table.metadata["common"]["section_path"] == ["Archive"]
    assert table.metadata["spreadsheet"] == {
        "sheet_name": "Archive",
        "sheet_state": "hidden",
        "cell_range": "A1:B2",
        "header_range": "A1:B1",
        "data_range": "A2:B2",
    }


def test_xlsx_backend_rejects_invalid_archives_and_excessive_cell_counts():
    from rag_document_parser import XlsxBackend

    with pytest.raises(ValueError, match="valid XLSX archive"):
        XlsxBackend().parse(b"not a ZIP archive", ".xlsx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["A", "B"])
    sheet.append([1, 2])

    with pytest.raises(ValueError, match=r"too many (?:occupied )?cells"):
        XlsxBackend(max_cells=3).parse(_workbook_bytes(workbook), ".xlsx")


def test_importing_xlsx_backend_does_not_eagerly_import_openpyxl():
    completed = subprocess.run(
        [
            sys.executable,
            "-c",
            (
                "import sys; "
                "from rag_document_parser import XlsxBackend; "
                "assert XlsxBackend.supported_suffixes == ('.xlsx',); "
                "assert 'openpyxl' not in sys.modules"
            ),
        ],
        check=False,
        capture_output=True,
        text=True,
    )

    assert completed.returncode == 0, completed.stderr


def test_xlsx_backend_uses_first_unstyled_row_as_header_before_mixed_data():
    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Name", "Tier"])
    sheet.append(["Alice", "Gold"])
    sheet.append(["Bob", 3])

    parsed = XlsxBackend().parse(_workbook_bytes(workbook), ".xlsx")

    table = _only_table(parsed)
    assert table.metadata["spreadsheet"]["header_range"] == "A1:B1"
    assert [column["text"] for column in table.content["columns"]] == [
        "Name",
        "Tier",
    ]
    assert [
        [cell["text"] for cell in row["cells"]] for row in table.content["rows"]
    ] == [
        ["Alice", "Gold"],
        ["Bob", "3"],
    ]


def test_xlsx_backend_keeps_single_full_width_merge_as_table_title():
    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    sheet = workbook.active
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "Sales"
    sheet.append(["January", "February", "March"])
    sheet.append([10, 20, 30])

    parsed = XlsxBackend().parse(_workbook_bytes(workbook), ".xlsx")

    assert [unit.type for unit in parsed.units] == ["text", "table"]
    assert parsed.units[0].content == "Sales"
    table = parsed.units[1]
    assert table.metadata["spreadsheet"]["header_range"] == "A2:C2"
    assert [column["text"] for column in table.content["columns"]] == [
        "January",
        "February",
        "March",
    ]
    assert [row["source_row"] for row in table.content["header_rows"]] == [2]
    assert [cell["text"] for cell in table.content["rows"][0]["cells"]] == [
        "10",
        "20",
        "30",
    ]


def test_xlsx_backend_expands_grouped_hidden_columns_into_metadata_and_cells():
    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["A", "B", "C", "D", "E"])
    sheet.append([1, 2, 3, 4, 5])
    sheet.column_dimensions.group("B", "D", hidden=True)

    parsed = XlsxBackend().parse(_workbook_bytes(workbook), ".xlsx")

    table = _only_table(parsed)
    assert table.metadata["spreadsheet"]["hidden_columns"] == ["B", "C", "D"]
    for row in [*table.content["header_rows"], *table.content["rows"]]:
        assert [cell.get("hidden", False) for cell in row["cells"]] == [
            False,
            True,
            True,
            True,
            False,
        ]
        assert "hidden" not in row["cells"][0]
        assert "hidden" not in row["cells"][4]


def test_xlsx_backend_serializes_data_table_formula_deterministically():
    from openpyxl.worksheet.formula import DataTableFormula

    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Scenario"
    sheet.append(["Item", "Input", "Output"])
    sheet.append(["A", 1, None])
    sheet["C2"] = DataTableFormula(
        ref="C2:C3",
        r1="B2",
        ca=True,
    )
    sheet.append(["B", 2, None])

    parsed = XlsxBackend().parse(_workbook_bytes(workbook), ".xlsx")

    table = _only_table(parsed)
    formula_cell = table.content["rows"][0]["cells"][2]
    expected = "=DATA_TABLE(ref=C2:C3, r1=B2, ca=TRUE)"
    assert formula_cell["formula"] == expected
    assert formula_cell["text"] == expected
    assert "object at" not in formula_cell["formula"]
    assert parsed.quality_warnings[0]["formula"] == expected


def test_xlsx_backend_warns_for_image_only_worksheet():
    from openpyxl.drawing.image import Image

    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Image only"
    sheet.add_image(Image(io.BytesIO(PNG_BYTES)), "A1")

    parsed = XlsxBackend().parse(_workbook_bytes(workbook), ".xlsx")

    assert parsed.units == []
    warnings = [
        warning
        for warning in parsed.quality_warnings
        if warning["type"] == "xlsx_images_unsupported"
    ]
    assert len(warnings) == 1
    assert warnings[0]["sheet_name"] == "Image only"
    assert warnings[0]["count"] == 1


def test_xlsx_backend_warns_for_chart_sheet():
    from openpyxl.chart import BarChart, Reference

    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Chart data"
    sheet.append(["Month", "Sales"])
    sheet.append(["January", 10])
    chart = BarChart()
    chart.add_data(
        Reference(sheet, min_col=2, min_row=1, max_row=2), titles_from_data=True
    )
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=2))
    chart_sheet = workbook.create_chartsheet("Dashboard")
    chart_sheet.add_chart(chart)

    parsed = XlsxBackend().parse(_workbook_bytes(workbook), ".xlsx")

    warnings = [
        warning
        for warning in parsed.quality_warnings
        if warning["type"] == "xlsx_charts_unsupported"
    ]
    assert len(warnings) == 1
    assert warnings[0]["sheet_name"] == "Dashboard"
    assert warnings[0]["chart_sheet"] is True
    assert warnings[0]["count"] == 1


def test_xlsx_backend_can_skip_hidden_sheets_with_warning():
    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    workbook.active.title = "Visible"
    hidden = workbook.create_sheet("Archive")
    hidden.sheet_state = "hidden"
    hidden.append(["Item", "Value"])
    hidden.append(["Old", 9])

    parsed = XlsxBackend(include_hidden_sheets=False).parse(
        _workbook_bytes(workbook),
        ".xlsx",
    )

    assert parsed.units == []
    assert len(parsed.quality_warnings) == 1
    warning = parsed.quality_warnings[0]
    assert {
        key: warning[key] for key in ("type", "severity", "sheet_name", "sheet_state")
    } == {
        "type": "xlsx_hidden_sheet_skipped",
        "severity": "low",
        "sheet_name": "Archive",
        "sheet_state": "hidden",
    }


def test_xlsx_backend_keeps_adjacent_declared_excel_tables_separate():
    from openpyxl.worksheet.table import Table

    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["First", "Value"])
    sheet.append(["A", 1])
    sheet.append(["Second", "Amount"])
    sheet.append(["B", 2])
    sheet.add_table(Table(displayName="FirstTable", ref="A1:B2"))
    sheet.add_table(Table(displayName="SecondTable", ref="A3:B4"))

    parsed = XlsxBackend().parse(_workbook_bytes(workbook), ".xlsx")

    tables = [unit for unit in parsed.units if unit.type == "table"]
    assert [table.metadata["spreadsheet"]["cell_range"] for table in tables] == [
        "A1:B2",
        "A3:B4",
    ]
    assert [table.metadata["table"]["row_count"] for table in tables] == [1, 1]
    assert [
        [column["text"] for column in table.content["columns"]] for table in tables
    ] == [["First", "Value"], ["Second", "Amount"]]


def test_xlsx_backend_uses_declared_tables_as_inference_barriers():
    from openpyxl.worksheet.table import Table

    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Top A", "Top B", "Top C", "Top D"])
    sheet.append(["Left 1", "Header 1", "Header 2", "Right 1"])
    sheet.append(["Left 2", "value", 1, "Right 2"])
    sheet.add_table(Table(displayName="InnerTable", ref="B2:C3"))

    parsed = XlsxBackend().parse(_workbook_bytes(workbook), ".xlsx")

    tables = [unit for unit in parsed.units if unit.type == "table"]
    assert len(tables) == 1
    assert tables[0].metadata["spreadsheet"]["cell_range"] == "B2:C3"
    assert [column["text"] for column in tables[0].content["columns"]] == [
        "Header 1",
        "Header 2",
    ]
    assert [cell["text"] for cell in tables[0].content["rows"][0]["cells"]] == [
        "value",
        "1",
    ]

    surrounding_text = "\n".join(
        str(unit.content) for unit in parsed.units if unit.type == "text"
    )
    assert "Top A | Top B | Top C | Top D" in surrounding_text
    assert "Left 1" in surrounding_text
    assert "Left 2" in surrounding_text
    assert "Right 1" in surrounding_text
    assert "Right 2" in surrounding_text
    assert "Header 1" not in surrounding_text
    assert "Header 2" not in surrounding_text
    assert "value" not in surrounding_text


def test_xlsx_backend_does_not_split_inferred_region_for_distant_table():
    from openpyxl.worksheet.table import Table

    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Name", "Value"])
    sheet.append(["A", 1])
    sheet.append(["B", 2, None, "Code", "Amount"])
    sheet.append(["C", 3, None, "X", 9])
    sheet.add_table(Table(displayName="DistantTable", ref="D3:E4"))

    parsed = XlsxBackend().parse(_workbook_bytes(workbook), ".xlsx")

    tables = [unit for unit in parsed.units if unit.type == "table"]
    assert [table.metadata["spreadsheet"]["cell_range"] for table in tables] == [
        "A1:B4",
        "D3:E4",
    ]
    assert [table.metadata["table"]["row_count"] for table in tables] == [3, 1]


def test_xlsx_backend_preserves_headerless_declared_table_rows():
    from openpyxl.worksheet.table import Table

    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["alpha", 1])
    sheet.append(["beta", 2])
    table = Table(displayName="HeaderlessTable", ref="A1:B2")
    table.headerRowCount = 0
    table._initialise_columns()
    table.tableColumns[0].name = "Name"
    table.tableColumns[1].name = "Value"
    sheet.add_table(table)

    parsed = XlsxBackend().parse(_workbook_bytes(workbook), ".xlsx")

    parsed_table = _only_table(parsed)
    assert [column["text"] for column in parsed_table.content["columns"]] == [
        "Name",
        "Value",
    ]
    assert parsed_table.content["header_rows"] == []
    assert [row["source_row"] for row in parsed_table.content["rows"]] == [1, 2]
    assert parsed_table.metadata["table"]["row_count"] == 2
    assert parsed_table.metadata["spreadsheet"] == {
        "sheet_name": "Sheet",
        "sheet_state": "visible",
        "cell_range": "A1:B2",
        "header_range": None,
        "data_range": "A1:B2",
    }


def test_xlsx_backend_validates_and_emits_empty_declared_table_footprint():
    from openpyxl.worksheet.table import Table

    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    sheet = workbook.active
    table = Table(displayName="EmptyTable", ref="A1:B2")
    table.headerRowCount = 0
    table._initialise_columns()
    table.tableColumns[0].name = "Name"
    table.tableColumns[1].name = "Value"
    sheet.add_table(table)
    data = _workbook_bytes(workbook)

    parsed_table = _only_table(XlsxBackend().parse(data, ".xlsx"))
    assert parsed_table.metadata["spreadsheet"]["cell_range"] == "A1:B2"
    assert parsed_table.metadata["table"]["row_count"] == 2
    assert [
        [cell["text"] for cell in row["cells"]] for row in parsed_table.content["rows"]
    ] == [["", ""], ["", ""]]

    with pytest.raises(ValueError, match="declared table ranges"):
        XlsxBackend(max_cells=3).parse(data, ".xlsx")


def test_xlsx_backend_rejects_reversed_declared_table_range():
    from openpyxl.worksheet.table import Table

    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    sheet = workbook.active
    table = Table(displayName="ReversedTable", ref="C1:A2")
    table.headerRowCount = 0
    sheet.add_table(table)

    with pytest.raises(ValueError, match="Invalid Excel table range"):
        XlsxBackend().parse(_workbook_bytes(workbook), ".xlsx")


def test_xlsx_preflight_follows_custom_worksheet_relationship_targets():
    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["A", "B"])
    sheet.append([1, 2])

    def rewrite(name: str, payload: bytes) -> tuple[str, bytes]:
        if name == "xl/worksheets/sheet1.xml":
            return "xl/custom/sheet1.xml", payload
        if name == "xl/_rels/workbook.xml.rels":
            payload = payload.replace(
                b"worksheets/sheet1.xml",
                b"custom/sheet1.xml",
            )
        if name == "[Content_Types].xml":
            payload = payload.replace(
                b"/xl/worksheets/sheet1.xml",
                b"/xl/custom/sheet1.xml",
            )
        return name, payload

    rewritten = _rewrite_xlsx(_workbook_bytes(workbook), rewrite)

    with pytest.raises(ValueError, match="too many cells before loading"):
        XlsxBackend(max_cells=1).parse(rewritten, ".xlsx")


def test_xlsx_preflight_rejects_columns_outside_excel_bounds():
    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "value"
    sheet.column_dimensions.group("B", "D", hidden=True)

    def rewrite(name: str, payload: bytes) -> tuple[str, bytes]:
        if name == "xl/worksheets/sheet1.xml":
            payload = payload.replace(b'max="4"', b'max="1000000000"')
        return name, payload

    rewritten = _rewrite_xlsx(_workbook_bytes(workbook), rewrite)

    with pytest.raises(ValueError, match="outside Excel bounds"):
        XlsxBackend(max_cells=10).parse(rewritten, ".xlsx")


def test_xlsx_preflight_caps_relationship_nodes_inside_xml():
    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    workbook.active["A1"] = "value"

    def rewrite(name: str, payload: bytes) -> tuple[str, bytes]:
        if name == "xl/_rels/workbook.xml.rels":
            relationships = b"".join(
                (
                    b'<Relationship Id="dummy%d" '
                    b'Type="urn:example:unused" Target="unused%d.xml"/>'
                )
                % (index, index)
                for index in range(30)
            )
            payload = payload.replace(
                b"</Relationships>", relationships + b"</Relationships>"
            )
        return name, payload

    rewritten = _rewrite_xlsx(_workbook_bytes(workbook), rewrite)

    with pytest.raises(ValueError, match="too many relationships"):
        XlsxBackend(max_archive_members=20).parse(rewritten, ".xlsx")


def test_xlsx_preflight_caps_sheet_nodes_inside_xml():
    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    workbook.create_sheet("Second")

    with pytest.raises(ValueError, match="too many sheets"):
        XlsxBackend(max_worksheets=1).parse(_workbook_bytes(workbook), ".xlsx")


def test_xlsx_preflight_uses_only_namespaced_sheet_relationship_ids(monkeypatch):
    import openpyxl

    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    workbook.active.append(["A", "B"])

    def rewrite(name: str, payload: bytes) -> tuple[str, bytes]:
        if name == "xl/workbook.xml":
            payload = payload.replace(
                b"<workbook xmlns:r=",
                (
                    b'<workbook xmlns:sr="http://purl.oclc.org/ooxml/'
                    b'officeDocument/relationships" xmlns:r='
                ),
            )
            payload = payload.replace(
                b"<sheet name=",
                b'<sheet id="decoy" sr:id="rId1" name=',
            )
        return name, payload

    def fail_if_loaded(*args, **kwargs):
        raise AssertionError("openpyxl must not run before the cell limit check")

    monkeypatch.setattr(openpyxl, "load_workbook", fail_if_loaded)
    rewritten = _rewrite_xlsx(_workbook_bytes(workbook), rewrite)

    with pytest.raises(ValueError, match="too many cells before loading"):
        XlsxBackend(max_cells=1).parse(rewritten, ".xlsx")


def test_xlsx_preflight_rejects_unknown_sheet_relationship_type_before_load(
    monkeypatch,
):
    import openpyxl

    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    workbook.active.append(["A", "B"])

    def rewrite(name: str, payload: bytes) -> tuple[str, bytes]:
        if name == "xl/_rels/workbook.xml.rels":
            payload = payload.replace(b'/worksheet"', b'/not-a-worksheet"')
        return name, payload

    def fail_if_loaded(*args, **kwargs):
        raise AssertionError("unexpected sheet types must fail before openpyxl")

    monkeypatch.setattr(openpyxl, "load_workbook", fail_if_loaded)
    rewritten = _rewrite_xlsx(_workbook_bytes(workbook), rewrite)

    with pytest.raises(ValueError, match="Unsupported workbook sheet relationship"):
        XlsxBackend(max_cells=1).parse(rewritten, ".xlsx")


def test_xlsx_preflight_rejects_reused_sheet_relationship_id_before_load(monkeypatch):
    import openpyxl

    from rag_document_parser import XlsxBackend

    workbook = Workbook()

    def rewrite(name: str, payload: bytes) -> tuple[str, bytes]:
        if name == "xl/workbook.xml":
            needle = b'<sheet name="Sheet" sheetId="1" state="visible" r:id="rId1" />'
            clone = b'<sheet name="Clone" sheetId="2" state="visible" r:id="rId1" />'
            payload = payload.replace(needle, needle + clone)
        return name, payload

    def fail_if_loaded(*args, **kwargs):
        raise AssertionError("reused sheet relationships must fail before openpyxl")

    monkeypatch.setattr(openpyxl, "load_workbook", fail_if_loaded)
    rewritten = _rewrite_xlsx(_workbook_bytes(workbook), rewrite)

    with pytest.raises(ValueError, match="same relationship identifier"):
        XlsxBackend().parse(rewritten, ".xlsx")


def test_xlsx_preflight_rejects_reused_worksheet_target_before_load(monkeypatch):
    import openpyxl

    from rag_document_parser import XlsxBackend

    workbook = Workbook()

    def rewrite(name: str, payload: bytes) -> tuple[str, bytes]:
        if name == "xl/workbook.xml":
            needle = b'<sheet name="Sheet" sheetId="1" state="visible" r:id="rId1" />'
            clone = b'<sheet name="Clone" sheetId="2" state="visible" r:id="rId99" />'
            payload = payload.replace(needle, needle + clone)
        if name == "xl/_rels/workbook.xml.rels":
            clone_relationship = (
                b'<Relationship Type="http://schemas.openxmlformats.org/'
                b'officeDocument/2006/relationships/worksheet" '
                b'Target="/xl/worksheets/sheet1.xml" Id="rId99" />'
            )
            payload = payload.replace(
                b"</Relationships>", clone_relationship + b"</Relationships>"
            )
        return name, payload

    def fail_if_loaded(*args, **kwargs):
        raise AssertionError("reused worksheet targets must fail before openpyxl")

    monkeypatch.setattr(openpyxl, "load_workbook", fail_if_loaded)
    rewritten = _rewrite_xlsx(_workbook_bytes(workbook), rewrite)

    with pytest.raises(ValueError, match="same package part"):
        XlsxBackend().parse(rewritten, ".xlsx")


def test_xlsx_preflight_counts_hyperlink_range_before_load(monkeypatch):
    import openpyxl

    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    workbook.active["A1"] = "link"
    workbook.active["A1"].hyperlink = "https://example.test"

    def rewrite(name: str, payload: bytes) -> tuple[str, bytes]:
        if name == "xl/worksheets/sheet1.xml":
            payload = payload.replace(
                b'relationships" ref="A1"',
                b'relationships" ref="A1:XFD1048576"',
            )
        return name, payload

    def fail_if_loaded(*args, **kwargs):
        raise AssertionError("hyperlink ranges must be bounded before openpyxl")

    monkeypatch.setattr(openpyxl, "load_workbook", fail_if_loaded)
    rewritten = _rewrite_xlsx(_workbook_bytes(workbook), rewrite)

    with pytest.raises(ValueError, match="too many cells before loading"):
        XlsxBackend().parse(rewritten, ".xlsx")


def test_xlsx_preflight_validates_comments_in_custom_worksheet_path(monkeypatch):
    import openpyxl
    from openpyxl.comments import Comment

    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    workbook.active["A1"].comment = Comment("note", "author")

    def rewrite(name: str, payload: bytes) -> tuple[str, bytes]:
        if name == "xl/worksheets/sheet1.xml":
            name = "xl/custom/sheet1.xml"
        elif name == "xl/worksheets/_rels/sheet1.xml.rels":
            name = "xl/custom/_rels/sheet1.xml.rels"
        elif name == "xl/_rels/workbook.xml.rels":
            payload = payload.replace(
                b"worksheets/sheet1.xml",
                b"custom/sheet1.xml",
            )
        elif name == "[Content_Types].xml":
            payload = payload.replace(
                b"/xl/worksheets/sheet1.xml",
                b"/xl/custom/sheet1.xml",
            )
        elif "comments" in name and name.endswith(".xml"):
            payload = payload.replace(
                b'ref="A1"',
                b'ref="A1:XFD1048576"',
            )
        return name, payload

    def fail_if_loaded(*args, **kwargs):
        raise AssertionError("comment references must be checked before openpyxl")

    monkeypatch.setattr(openpyxl, "load_workbook", fail_if_loaded)
    rewritten = _rewrite_xlsx(_workbook_bytes(workbook), rewrite)

    with pytest.raises(ValueError, match="Invalid worksheet cell reference"):
        XlsxBackend().parse(rewritten, ".xlsx")


def test_xlsx_preflight_caps_worksheet_relationship_nodes(monkeypatch):
    import openpyxl

    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    workbook.active["A1"] = "link"
    workbook.active["A1"].hyperlink = "https://example.test"

    def rewrite(name: str, payload: bytes) -> tuple[str, bytes]:
        if name == "xl/worksheets/_rels/sheet1.xml.rels":
            relationships = b"".join(
                (
                    b'<Relationship Id="dummy%d" '
                    b'Type="urn:example:unused" Target="unused%d.xml"/>'
                )
                % (index, index)
                for index in range(30)
            )
            payload = payload.replace(
                b"</Relationships>", relationships + b"</Relationships>"
            )
        return name, payload

    def fail_if_loaded(*args, **kwargs):
        raise AssertionError("worksheet relationships must be capped before openpyxl")

    monkeypatch.setattr(openpyxl, "load_workbook", fail_if_loaded)
    rewritten = _rewrite_xlsx(_workbook_bytes(workbook), rewrite)

    with pytest.raises(ValueError, match="worksheets have too many relationships"):
        XlsxBackend(max_archive_members=20).parse(rewritten, ".xlsx")


def test_xlsx_preflight_rejects_non_cell_row_children_before_load(monkeypatch):
    import openpyxl

    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    workbook.active["A1"] = 1

    def rewrite(name: str, payload: bytes) -> tuple[str, bytes]:
        if name == "xl/worksheets/sheet1.xml":
            payload = payload.replace(b'<c r="A1"', b'<foo r="A1"')
            payload = payload.replace(b"</c>", b"</foo>", 1)
        return name, payload

    def fail_if_loaded(*args, **kwargs):
        raise AssertionError("unexpected row children must fail before openpyxl")

    monkeypatch.setattr(openpyxl, "load_workbook", fail_if_loaded)
    rewritten = _rewrite_xlsx(_workbook_bytes(workbook), rewrite)

    with pytest.raises(ValueError, match="row contains an unexpected element"):
        XlsxBackend().parse(rewritten, ".xlsx")


def test_xlsx_preflight_rejects_nested_rows_before_load(monkeypatch):
    import openpyxl

    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    workbook.active["A1"] = 1

    def rewrite(name: str, payload: bytes) -> tuple[str, bytes]:
        if name == "xl/worksheets/sheet1.xml":
            payload = payload.replace(
                b'<c r="A1" t="n"><v>1</v></c>',
                b'<row><c r="A1" t="n"><v>1</v></c></row>',
            )
        return name, payload

    def fail_if_loaded(*args, **kwargs):
        raise AssertionError("nested rows must fail before openpyxl")

    monkeypatch.setattr(openpyxl, "load_workbook", fail_if_loaded)
    rewritten = _rewrite_xlsx(_workbook_bytes(workbook), rewrite)

    with pytest.raises(ValueError, match="rows must be direct children"):
        XlsxBackend().parse(rewritten, ".xlsx")


def test_xlsx_preflight_bounds_inferred_cell_columns_before_load(monkeypatch):
    import openpyxl

    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    workbook.active["A1"] = 1
    workbook.active["XFD1"] = 2

    def rewrite(name: str, payload: bytes) -> tuple[str, bytes]:
        if name == "xl/worksheets/sheet1.xml":
            payload = payload.replace(
                b"</row>",
                b'<c t="n"><v>3</v></c></row>',
                1,
            )
        return name, payload

    def fail_if_loaded(*args, **kwargs):
        raise AssertionError("inferred cell columns must be checked before openpyxl")

    monkeypatch.setattr(openpyxl, "load_workbook", fail_if_loaded)
    rewritten = _rewrite_xlsx(_workbook_bytes(workbook), rewrite)

    with pytest.raises(ValueError, match="Inferred worksheet cell is outside"):
        XlsxBackend().parse(rewritten, ".xlsx")


def test_xlsx_preflight_bounds_inferred_rows_before_load(monkeypatch):
    import openpyxl

    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    workbook.active["A1048576"] = 1

    def rewrite(name: str, payload: bytes) -> tuple[str, bytes]:
        if name == "xl/worksheets/sheet1.xml":
            payload = payload.replace(
                b"</sheetData>",
                b'<row><c t="n"><v>2</v></c></row></sheetData>',
            )
        return name, payload

    def fail_if_loaded(*args, **kwargs):
        raise AssertionError("inferred rows must be checked before openpyxl")

    monkeypatch.setattr(openpyxl, "load_workbook", fail_if_loaded)
    rewritten = _rewrite_xlsx(_workbook_bytes(workbook), rewrite)

    with pytest.raises(ValueError, match="Inferred worksheet row is outside"):
        XlsxBackend().parse(rewritten, ".xlsx")


@pytest.mark.parametrize(
    ("mutation", "message"),
    [
        ("duplicate", "exactly one standard workbook part"),
        ("child_override", "entries must not contain child elements"),
    ],
)
def test_xlsx_preflight_rejects_ambiguous_workbook_manifest_before_load(
    monkeypatch,
    mutation,
    message,
):
    import openpyxl

    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    workbook_type = (
        b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
    )
    override = (
        b'<Override PartName="/xl/workbook.xml" ContentType="' + workbook_type + b'" />'
    )

    def rewrite(name: str, payload: bytes) -> tuple[str, bytes]:
        if name != "[Content_Types].xml":
            return name, payload
        if mutation == "duplicate":
            payload = payload.replace(b"</Types>", override + b"</Types>")
        else:
            payload = payload.replace(
                override,
                (
                    b'<Override PartName="/xl/workbook.xml" ContentType="'
                    + workbook_type
                    + b'"><PartName>/xl/evilworkbook.xml</PartName></Override>'
                ),
            )
        return name, payload

    def fail_if_loaded(*args, **kwargs):
        raise AssertionError("ambiguous workbook parts must fail before openpyxl")

    monkeypatch.setattr(openpyxl, "load_workbook", fail_if_loaded)
    rewritten = _rewrite_xlsx(_workbook_bytes(workbook), rewrite)

    with pytest.raises(ValueError, match=message):
        XlsxBackend().parse(rewritten, ".xlsx")


@pytest.mark.parametrize(
    ("mutation", "message"),
    [
        ("tag", "sheets contain an unexpected element"),
        ("child_override", "sheet entries must be empty"),
    ],
)
def test_xlsx_preflight_rejects_sheet_descriptor_overrides_before_load(
    monkeypatch,
    mutation,
    message,
):
    import openpyxl

    from rag_document_parser import XlsxBackend

    workbook = Workbook()

    def rewrite(name: str, payload: bytes) -> tuple[str, bytes]:
        if name != "xl/workbook.xml":
            return name, payload
        if mutation == "tag":
            payload = payload.replace(b"<sheet ", b"<foo ", 1)
        else:
            payload = payload.replace(
                b'r:id="rId1" />',
                b'r:id="rId1"><id>rId2</id></sheet>',
                1,
            )
        return name, payload

    def fail_if_loaded(*args, **kwargs):
        raise AssertionError("sheet descriptor overrides must fail before openpyxl")

    monkeypatch.setattr(openpyxl, "load_workbook", fail_if_loaded)
    rewritten = _rewrite_xlsx(_workbook_bytes(workbook), rewrite)

    with pytest.raises(ValueError, match=message):
        XlsxBackend().parse(rewritten, ".xlsx")


@pytest.mark.parametrize(
    ("part_kind", "mutation", "message"),
    [
        ("workbook", "tag", "relationships contain an unexpected element"),
        ("workbook", "child_override", "relationship entries must be empty"),
        ("worksheet", "tag", "relationships contain an unexpected element"),
        ("worksheet", "child_override", "relationship entries must be empty"),
    ],
)
def test_xlsx_preflight_rejects_relationship_descriptor_overrides_before_load(
    monkeypatch,
    part_kind,
    mutation,
    message,
):
    import openpyxl
    from openpyxl.comments import Comment

    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    if part_kind == "worksheet":
        workbook.active["A1"].comment = Comment("note", "author")
        relationship_part = "xl/worksheets/_rels/sheet1.xml.rels"
        relationship_id = b"comments"
    else:
        relationship_part = "xl/_rels/workbook.xml.rels"
        relationship_id = b"rId1"

    def rewrite(name: str, payload: bytes) -> tuple[str, bytes]:
        if name != relationship_part:
            return name, payload
        if mutation == "tag":
            payload = payload.replace(b"<Relationship ", b"<foo ", 1)
        else:
            closing = b'Id="' + relationship_id + b'" />'
            payload = payload.replace(
                closing,
                (
                    b'Id="'
                    + relationship_id
                    + b'"><Target>/xl/evil.xml</Target></Relationship>'
                ),
                1,
            )
        return name, payload

    def fail_if_loaded(*args, **kwargs):
        raise AssertionError("relationship overrides must fail before openpyxl")

    monkeypatch.setattr(openpyxl, "load_workbook", fail_if_loaded)
    rewritten = _rewrite_xlsx(_workbook_bytes(workbook), rewrite)

    with pytest.raises(ValueError, match=message):
        XlsxBackend().parse(rewritten, ".xlsx")


@pytest.mark.parametrize(
    ("target", "message"),
    [
        (b"worksheets\\sheet1.xml", "contains a backslash"),
        (b"/../xl/workbook.xml", "not a canonical package path"),
    ],
)
def test_xlsx_preflight_rejects_ambiguous_relationship_targets_before_load(
    monkeypatch,
    target,
    message,
):
    import openpyxl

    from rag_document_parser import XlsxBackend

    workbook = Workbook()

    def rewrite(name: str, payload: bytes) -> tuple[str, bytes]:
        if name == "xl/_rels/workbook.xml.rels":
            payload = payload.replace(
                b"/xl/worksheets/sheet1.xml",
                target,
                1,
            )
        return name, payload

    def fail_if_loaded(*args, **kwargs):
        raise AssertionError("ambiguous targets must fail before openpyxl")

    monkeypatch.setattr(openpyxl, "load_workbook", fail_if_loaded)
    rewritten = _rewrite_xlsx(_workbook_bytes(workbook), rewrite)

    with pytest.raises(ValueError, match=message):
        XlsxBackend().parse(rewritten, ".xlsx")


@pytest.mark.parametrize(
    ("mutation", "message"),
    [
        ("tag", "comment list contains an unexpected element"),
        ("child_ref", "cell references must be attributes"),
    ],
)
def test_xlsx_preflight_rejects_comment_descriptor_overrides_before_load(
    monkeypatch,
    mutation,
    message,
):
    import openpyxl
    from openpyxl.comments import Comment

    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    workbook.active["A1"].comment = Comment("note", "author")

    def rewrite(name: str, payload: bytes) -> tuple[str, bytes]:
        if "comments" not in name or not name.endswith(".xml"):
            return name, payload
        if mutation == "tag":
            payload = payload.replace(b"<comment ", b"<foo ", 1)
            payload = payload.replace(b"</comment>", b"</foo>", 1)
        else:
            payload = payload.replace(
                b'shapeId="0">',
                b'shapeId="0"><ref>A1:XFD1048576</ref>',
                1,
            )
        return name, payload

    def fail_if_loaded(*args, **kwargs):
        raise AssertionError("comment overrides must fail before openpyxl")

    monkeypatch.setattr(openpyxl, "load_workbook", fail_if_loaded)
    rewritten = _rewrite_xlsx(_workbook_bytes(workbook), rewrite)

    with pytest.raises(ValueError, match=message):
        XlsxBackend().parse(rewritten, ".xlsx")


def test_xlsx_preflight_rejects_hyperlink_ref_child_override_before_load(
    monkeypatch,
):
    import openpyxl

    from rag_document_parser import XlsxBackend

    workbook = Workbook()
    workbook.active["A1"] = "link"
    workbook.active["A1"].hyperlink = "https://example.test"

    def rewrite(name: str, payload: bytes) -> tuple[str, bytes]:
        if name == "xl/worksheets/sheet1.xml":
            payload = payload.replace(
                b'ref="A1" r:id="rId1" />',
                (b'ref="A1" r:id="rId1"><ref>A1:XFD1048576</ref></hyperlink>'),
                1,
            )
        return name, payload

    def fail_if_loaded(*args, **kwargs):
        raise AssertionError("hyperlink ref overrides must fail before openpyxl")

    monkeypatch.setattr(openpyxl, "load_workbook", fail_if_loaded)
    rewritten = _rewrite_xlsx(_workbook_bytes(workbook), rewrite)

    with pytest.raises(ValueError, match="hyperlink entries must be empty"):
        XlsxBackend().parse(rewritten, ".xlsx")
