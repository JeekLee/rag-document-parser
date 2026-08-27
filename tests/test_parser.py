from __future__ import annotations

import hashlib
import json


def _s3_config():
    from rag_document_parser import S3Config

    return S3Config(
        endpoint="http://minio.test",
        bucket="rag-assets",
        access_key="access",
        secret_key="secret",
        prefix="documents",
    )


class _RecordingAssetStore:
    def __init__(self) -> None:
        self.uploads: list[tuple[str, bytes, str]] = []

    def put(self, key: str, data: bytes, content_type: str) -> str:
        self.uploads.append((key, data, content_type))
        return f"s3://rag-assets/documents/{key}"


def test_parse_markdown_document_is_disabled_by_default_without_llm(monkeypatch):
    from rag_document_parser import RagDocumentParser

    def fail_chat_json(prompt, cfg):
        raise AssertionError("parse() must not call LLM enrichment")

    monkeypatch.setattr(
        "rag_document_parser.llm.chat_json",
        fail_chat_json,
        raising=False,
    )

    raw = (
        "# 요양급여 기준\n\n"
        "코로나19 대면투약관리료는 다음 기준에 따라 산정한다.\n\n"
        "| 대상 | 청구방법 |\n"
        "| --- | --- |\n"
        "| 약국 | 대면투약관리료 코드로 청구 |\n"
    ).encode()

    try:
        RagDocumentParser(object_storage=_s3_config()).parse(
            raw,
            suffix=".md",
            source_id="notice-1/attachment-1",
            source_name="기준.md",
        )
    except ValueError as exc:
        message = str(exc)
        assert "Unsupported format" in message
        assert ".md" in message
        assert ".txt" not in message
        assert ".markdown" not in message
    else:
        raise AssertionError("expected .md to be disabled by default")


def test_markdown_backend_returns_evidence_units():
    from rag_document_parser import EvidenceUnit
    from rag_document_parser.evidence_unit_extraction.formats.markdown import (
        MarkdownBackend,
    )

    raw = (
        "# Section\n\n"
        "Plain paragraph.\n\n"
        "| A | B |\n"
        "| --- | --- |\n"
        "| one | two |\n"
    ).encode()

    parsed = MarkdownBackend().parse(raw, ".md")

    assert not hasattr(parsed, "chunks")
    assert [unit.id for unit in parsed.units] == ["b1", "b2"]
    assert all(isinstance(unit, EvidenceUnit) for unit in parsed.units)

    text_unit, table_unit = parsed.units
    assert text_unit.type == "text"
    assert text_unit.source.text == "section: Section\nPlain paragraph."
    assert text_unit.format == "plain"
    assert text_unit.content == "Plain paragraph."
    assert text_unit.metadata["common"] == {
        "chunk_kind": "text",
        "section_path": ["Section"],
        "display_format": "plain",
    }

    assert table_unit.type == "table"
    assert table_unit.source.text == "section: Section\ncolumns: A | B\nrow 1: A=one; B=two"
    assert table_unit.format == "structured_table"
    assert table_unit.content["columns"] == [
        {"id": "c1", "text": "A"},
        {"id": "c2", "text": "B"},
    ]
    assert table_unit.metadata["common"]["display_format"] == "structured_table"


def test_parse_result_to_dict_is_json_serializable():
    from rag_document_parser import RagDocumentParser

    raw = b"<p>plain paragraph</p>"
    payload = RagDocumentParser(object_storage=_s3_config()).parse(
        raw, suffix=".html"
    ).to_dict()

    assert payload["source"]["sha256"] == hashlib.sha256(raw).hexdigest()
    assert payload["assets"] == []
    assert "chunks" not in payload
    assert "preview_markdown" not in payload
    assert "source_pointer" not in payload["units"][0]
    assert "embedding_text" not in payload["units"][0]
    assert "summary" not in payload["units"][0]
    assert payload["units"][0]["source"] == {
        "kind": "text",
        "text": "plain paragraph",
    }
    assert payload["units"][0]["format"] == "plain"
    assert payload["units"][0]["content"] == "plain paragraph"
    assert "evidence" not in payload["units"][0]
    assert json.loads(json.dumps(payload, ensure_ascii=False)) == payload


def test_parser_routes_uppercase_xlsx_and_preserves_source_metadata():
    from io import BytesIO

    from openpyxl import Workbook

    from rag_document_parser import RagDocumentParser

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Benefits"
    sheet.append(["Item", "Amount"])
    sheet.append(["Clinic", 1000])
    output = BytesIO()
    workbook.save(output)
    raw = output.getvalue()

    result = RagDocumentParser(object_storage=_s3_config()).parse(
        raw,
        suffix=".XLSX",
        source_id="benefits-2026",
        source_name="benefits.xlsx",
    )

    assert result.source.suffix == ".xlsx"
    assert result.source.sha256 == hashlib.sha256(raw).hexdigest()
    assert result.source.bytes == len(raw)
    assert result.source.id == "benefits-2026"
    assert result.source.name == "benefits.xlsx"
    assert result.assets == []
    assert [unit.type for unit in result.units] == ["table"]
    table = result.units[0]
    assert [column["text"] for column in table.content["columns"]] == [
        "Item",
        "Amount",
    ]
    assert table.content["rows"][0]["cells"][0]["text"] == "Clinic"
    assert table.metadata["spreadsheet"]["sheet_name"] == "Benefits"


def test_source_does_not_require_position_offsets():
    from rag_document_parser import RagDocumentParser

    raw = "<h1>H</h1><p>한글</p>".encode()
    unit = RagDocumentParser(object_storage=_s3_config()).parse(
        raw, suffix=".html"
    ).units[0]

    assert unit.source.text == "section: H\n한글"
    assert unit.source.to_dict() == {"kind": "text", "text": "section: H\n한글"}
    assert unit.metadata["common"]["section_path"] == ["H"]
    assert not hasattr(unit, "source_pointer")


def test_parser_allows_asset_free_documents_without_storage_config():
    from rag_document_parser import RagDocumentParser

    result = RagDocumentParser().parse(b"<p>text only</p>", suffix=".html")

    assert result.assets == []
    assert result.units[0].content == "text only"
    RagDocumentParser(object_storage=_s3_config())


def test_custom_backend_can_be_registered_for_suffix():
    from rag_document_parser import (
        EvidenceUnit,
        RagDocumentParser,
        SourceEvidence,
    )
    from rag_document_parser.evidence_unit_extraction.backend import ParsedDocument

    class CustomBackend:
        calls: list[tuple[bytes, str]]

        def __init__(self) -> None:
            self.calls = []

        def parse(self, data: bytes, suffix: str) -> ParsedDocument:
            self.calls.append((data, suffix))
            return ParsedDocument(
                units=[
                    EvidenceUnit(
                        id="c1",
                        type="text",
                        format="plain",
                        source=SourceEvidence(kind="text", text="custom text"),
                        content="custom text",
                    )
                ],
                quality_warnings=[
                    {
                        "type": "custom_warning",
                        "severity": "low",
                        "message": "backend warning",
                    }
                ],
            )

    backend = CustomBackend()
    result = RagDocumentParser(
        object_storage=_s3_config(),
        backends={".custom": backend},
    ).parse(b"custom bytes", suffix=".CUSTOM")

    assert backend.calls == [(b"custom bytes", ".custom")]
    assert result.source.suffix == ".custom"
    assert result.units[0].type == "text"
    assert result.units[0].format == "plain"
    assert result.units[0].content == "custom text"
    assert result.units[0].source.text == "custom text"
    assert result.quality_warnings == [
        {
            "type": "custom_warning",
            "severity": "low",
            "message": "backend warning",
        }
    ]


def test_unsupported_suffix_fails_before_llm_call(monkeypatch):
    from rag_document_parser import RagDocumentParser

    def fail_chat_json(prompt, cfg):
        raise AssertionError("parse() must not call LLM enrichment")

    monkeypatch.setattr(
        "rag_document_parser.llm.chat_json",
        fail_chat_json,
        raising=False,
    )

    parser = RagDocumentParser(object_storage=_s3_config())
    try:
        parser.parse(b"not supported", suffix=".docx")
    except ValueError as exc:
        assert "Unsupported format" in str(exc)
        assert ".docx" in str(exc)
    else:
        raise AssertionError("expected ValueError")


def test_image_assets_are_uploaded_to_asset_store_and_linked_in_evidence():
    from rag_document_parser import (
        EvidenceUnit,
        PendingAsset,
        RagDocumentParser,
        SourceEvidence,
    )
    from rag_document_parser.evidence_unit_extraction.backend import ParsedDocument

    asset_store = _RecordingAssetStore()

    class ImageBackend:
        def parse(self, data: bytes, suffix: str) -> ParsedDocument:
            return ParsedDocument(
                units=[
                    EvidenceUnit(
                        id="img-unit-1",
                        type="image",
                        format="asset_ref",
                        source=SourceEvidence(
                            kind="image",
                            text="summary: 요양급여 청구 절차를 보여주는 이미지",
                        ),
                        content={
                            "asset_id": "img-0001",
                            "caption": "청구 절차 이미지",
                        },
                        metadata={
                            "common": {
                                "chunk_kind": "image",
                                "display_format": "image",
                            }
                        },
                    )
                ],
                assets=[
                    PendingAsset(
                        id="img-0001",
                        kind="image",
                        data=b"png bytes",
                        mime="image/png",
                        ext="png",
                    )
                ],
            )

    raw = b"fake source document"
    try:
        RagDocumentParser(backends={".imgdoc": ImageBackend()}).parse(
            raw,
            suffix=".imgdoc",
        )
    except ValueError as exc:
        assert "asset_store is required" in str(exc)
    else:
        raise AssertionError("expected assets without a store to be rejected")

    result = RagDocumentParser(
        asset_store=asset_store,
        backends={".imgdoc": ImageBackend()},
    ).parse(raw, suffix=".imgdoc")

    document_hash = hashlib.sha256(raw).hexdigest()
    assert asset_store.uploads == [
        (
            f"{document_hash}/assets/img-0001.png",
            b"png bytes",
            "image/png",
        )
    ]
    assert result.assets[0].to_dict() == {
        "id": "img-0001",
        "kind": "image",
        "uri": f"s3://rag-assets/documents/{document_hash}/assets/img-0001.png",
        "mime": "image/png",
        "ext": "png",
        "sha256": hashlib.sha256(b"png bytes").hexdigest(),
        "bytes": len(b"png bytes"),
        "metadata": {},
    }
    assert result.units[0].content == {
        "asset_id": "img-0001",
        "caption": "청구 절차 이미지",
        "uri": f"s3://rag-assets/documents/{document_hash}/assets/img-0001.png",
        "mime": "image/png",
        "ext": "png",
        "sha256": hashlib.sha256(b"png bytes").hexdigest(),
        "bytes": len(b"png bytes"),
    }


def test_parser_registers_html_backend_and_uploads_nested_html_images():
    from rag_document_parser import RagDocumentParser

    asset_store = _RecordingAssetStore()

    import base64

    image_bytes = b"png bytes"
    data_uri = f"data:image/png;base64,{base64.b64encode(image_bytes).decode()}"
    raw = f"""
    <table>
      <tr><th>Item</th><th>Image</th></tr>
      <tr><td>Criteria</td><td><img src="{data_uri}" alt="cell chart"></td></tr>
    </table>
    """.encode()

    result = RagDocumentParser(asset_store=asset_store).parse(
        raw,
        suffix=".HTML",
    )
    document_hash = hashlib.sha256(raw).hexdigest()

    assert result.source.suffix == ".html"
    assert asset_store.uploads == [
        (
            f"{document_hash}/assets/img-0001.png",
            image_bytes,
            "image/png",
        )
    ]
    image_child = result.units[0].content["rows"][0]["cells"][1]["children"][0]
    assert image_child["content"]["uri"] == (
        f"s3://rag-assets/documents/{document_hash}/assets/img-0001.png"
    )
    assert image_child["content"]["sha256"] == hashlib.sha256(image_bytes).hexdigest()
