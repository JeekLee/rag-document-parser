from __future__ import annotations

import io
import math
import re
import warnings as python_warnings
from bisect import bisect_right
from collections import Counter, defaultdict
from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Any

from ....models import EvidenceUnit, ParsedDocument, SourceEvidence
from ...schema import common_metadata, structured_table
from .package_validation import (
    EXCEL_MAX_COLUMNS as _EXCEL_MAX_COLUMNS,
    EXCEL_MAX_ROWS as _EXCEL_MAX_ROWS,
    XlsxPackageValidator,
    column_number as _column_number,
)

_DEFAULT_MAX_UNCOMPRESSED_BYTES = 256 * 1024 * 1024
_DEFAULT_MAX_ARCHIVE_MEMBERS = 10_000
_DEFAULT_MAX_CELLS = 1_000_000
_DEFAULT_MAX_WORKSHEETS = 256
_GENERAL_NUMBER_FORMAT = "General"


@dataclass(frozen=True)
class _Region:
    min_row: int
    max_row: int
    min_col: int
    max_col: int

    @property
    def width(self) -> int:
        return self.max_col - self.min_col + 1

    @property
    def height(self) -> int:
        return self.max_row - self.min_row + 1

    @property
    def area(self) -> int:
        return self.width * self.height


@dataclass(frozen=True)
class _MergedRange:
    min_row: int
    max_row: int
    min_col: int
    max_col: int

    @property
    def anchor(self) -> tuple[int, int]:
        return self.min_row, self.min_col

    @property
    def area(self) -> int:
        return (self.max_row - self.min_row + 1) * (self.max_col - self.min_col + 1)


@dataclass(frozen=True)
class _DeclaredTable:
    region: _Region
    header_row_count: int
    column_labels: tuple[str, ...]


@dataclass
class _RegionIntervalNode:
    center: int
    by_start: list[_Region]
    by_end: list[_Region]
    starts: list[int]
    negative_ends: list[int]
    left: _RegionIntervalNode | None = None
    right: _RegionIntervalNode | None = None


@dataclass(frozen=True)
class _RegionSpatialIndex:
    rows: _RegionIntervalNode | None
    columns: _RegionIntervalNode | None
    row_starts: tuple[int, ...]
    row_ends: tuple[int, ...]
    column_starts: tuple[int, ...]
    column_ends: tuple[int, ...]


@dataclass(frozen=True)
class _RowProfile:
    total: int = 0
    strings: int = 0
    non_strings: int = 0
    formulas: int = 0
    bold: int = 0

    @property
    def bold_ratio(self) -> float:
        return self.bold / self.total if self.total else 0.0

    @property
    def text_ratio(self) -> float:
        return self.strings / self.total if self.total else 0.0


@dataclass
class _ParseState:
    quality_warnings: list[dict[str, Any]] = field(default_factory=list)
    block_index: int = 1
    table_index: int = 1
    declared_table_count: int = 0
    represented_cells: int = 0

    def next_block_id(self) -> str:
        block_id = f"b{self.block_index}"
        self.block_index += 1
        return block_id

    def next_table_id(self) -> str:
        table_id = f"t{self.table_index}"
        self.table_index += 1
        return table_id


@dataclass
class XlsxBackend:
    """Extract modern Excel workbooks without evaluating their formulas."""

    supported_suffixes = (".xlsx",)

    include_hidden_sheets: bool = True
    max_uncompressed_bytes: int = _DEFAULT_MAX_UNCOMPRESSED_BYTES
    max_archive_members: int = _DEFAULT_MAX_ARCHIVE_MEMBERS
    max_cells: int = _DEFAULT_MAX_CELLS
    max_worksheets: int = _DEFAULT_MAX_WORKSHEETS

    def parse(self, data: bytes, suffix: str) -> ParsedDocument:
        try:
            from openpyxl import load_workbook
        except (ImportError, ModuleNotFoundError) as exc:
            raise NotImplementedError(
                "XLSX extraction requires the optional 'openpyxl' dependency. "
                "Install rag-document-parser with the 'xlsx' extra."
            ) from exc

        XlsxPackageValidator(
            max_members=self.max_archive_members,
            max_uncompressed_bytes=self.max_uncompressed_bytes,
            max_cells=self.max_cells,
            max_worksheets=self.max_worksheets,
        ).validate(data)

        formula_workbook: Any | None = None
        value_workbook: Any | None = None
        reader_warnings: list[python_warnings.WarningMessage] = []
        try:
            with python_warnings.catch_warnings(record=True) as caught:
                python_warnings.simplefilter("always")
                formula_workbook = load_workbook(
                    io.BytesIO(data),
                    data_only=False,
                    read_only=False,
                    keep_links=False,
                    rich_text=True,
                )
                value_workbook = load_workbook(
                    io.BytesIO(data),
                    data_only=True,
                    read_only=False,
                    keep_links=False,
                    rich_text=True,
                )
                reader_warnings.extend(caught)
        except (MemoryError, RecursionError):
            raise
        except Exception as exc:
            raise ValueError("Invalid or unsupported XLSX workbook") from exc

        state = _ParseState()
        _append_reader_warnings(state, reader_warnings)
        units: list[EvidenceUnit] = []
        try:
            for chartsheet in formula_workbook.chartsheets:
                charts = list(getattr(chartsheet, "_charts", []))
                state.quality_warnings.append(
                    {
                        "type": "xlsx_charts_unsupported",
                        "severity": "medium",
                        "message": (
                            "Chart sheets are not extracted by the XLSX backend."
                        ),
                        "sheet_name": chartsheet.title,
                        "count": max(1, len(charts)),
                        "chart_sheet": True,
                    }
                )
            for worksheet in formula_workbook.worksheets:
                if (
                    worksheet.sheet_state != "visible"
                    and not self.include_hidden_sheets
                ):
                    state.quality_warnings.append(
                        {
                            "type": "xlsx_hidden_sheet_skipped",
                            "severity": "low",
                            "message": "A hidden worksheet was skipped by configuration.",
                            "sheet_name": worksheet.title,
                            "sheet_state": worksheet.sheet_state,
                        }
                    )
                    continue

                value_worksheet = value_workbook[worksheet.title]
                _parse_worksheet(
                    worksheet,
                    value_worksheet,
                    units,
                    state,
                    max_cells=self.max_cells,
                    max_tables=self.max_archive_members,
                )
        finally:
            if formula_workbook is not None:
                formula_workbook.close()
            if value_workbook is not None:
                value_workbook.close()

        return ParsedDocument(
            units=units,
            quality_warnings=state.quality_warnings,
        )


class _SheetReader:
    def __init__(
        self,
        formula_worksheet: Any,
        value_worksheet: Any,
        state: _ParseState,
    ) -> None:
        self.formula_worksheet = formula_worksheet
        self.value_worksheet = value_worksheet
        self.state = state
        self._display_cache: dict[tuple[int, int], str] = {}
        self._warned_formula_cells: set[str] = set()
        self.hidden_columns = _hidden_column_numbers(formula_worksheet)

    def cell(self, row: int, column: int) -> Any:
        return self.formula_worksheet.cell(row=row, column=column)

    def text(self, row: int, column: int) -> str:
        key = (row, column)
        if key in self._display_cache:
            return self._display_cache[key]

        formula_cell = self.cell(row, column)
        value = formula_cell.value
        if formula_cell.data_type == "f":
            formula = _formula_text(value)
            cached_cell = self.value_worksheet.cell(row=row, column=column)
            if cached_cell.value is None:
                value = formula
                self._warn_missing_formula_cache(formula_cell.coordinate, formula)
            else:
                value = cached_cell.value

        display = _format_cell_value(value, formula_cell.number_format)
        self._display_cache[key] = display
        return display

    def cell_payload(
        self,
        row: int,
        column: int,
        *,
        column_id: str,
        rowspan: int = 1,
        colspan: int = 1,
    ) -> dict[str, Any]:
        cell = self.cell(row, column)
        children: list[dict[str, Any]] = []
        comment = getattr(cell, "comment", None)
        if comment is not None and str(comment.text or "").strip():
            metadata: dict[str, Any] = {"role": "cell_comment"}
            if comment.author:
                metadata["author"] = str(comment.author)
            children.append(
                {
                    "type": "text",
                    "format": "plain",
                    "content": _clean_text(str(comment.text)),
                    "metadata": metadata,
                }
            )

        payload: dict[str, Any] = {
            "column_id": column_id,
            "text": self.text(row, column),
            "rowspan": rowspan,
            "colspan": colspan,
            "children": children,
            "address": cell.coordinate,
        }
        if cell.data_type == "f":
            payload["formula"] = _formula_text(cell.value)
        number_format = str(cell.number_format or _GENERAL_NUMBER_FORMAT)
        if number_format != _GENERAL_NUMBER_FORMAT:
            payload["number_format"] = number_format
        hyperlink = getattr(cell, "hyperlink", None)
        if hyperlink is not None:
            target = str(hyperlink.target or hyperlink.location or "").strip()
            if target:
                payload["hyperlink"] = target
        if column in self.hidden_columns:
            payload["hidden"] = True
        return payload

    def _warn_missing_formula_cache(self, coordinate: str, formula: str) -> None:
        if coordinate in self._warned_formula_cells:
            return
        self._warned_formula_cells.add(coordinate)
        self.state.quality_warnings.append(
            {
                "type": "xlsx_formula_cache_missing",
                "severity": "medium",
                "message": (
                    "The workbook did not contain a cached result for a formula; "
                    "the original formula was preserved as text."
                ),
                "sheet_name": self.formula_worksheet.title,
                "cell": coordinate,
                "formula": formula,
            }
        )


def _cell_payload_source_value(payload: dict[str, Any]) -> str:
    text = str(payload["text"])
    hyperlink = str(payload.get("hyperlink", "")).strip()
    if hyperlink and hyperlink != text:
        text = f"{text} ({hyperlink})" if text else hyperlink
    child_texts = [
        str(child.get("content", "")).strip()
        for child in payload["children"]
        if str(child.get("content", "")).strip()
    ]
    if child_texts:
        comment_text = " / ".join(child_texts)
        text = f"{text} [comment: {comment_text}]" if text else comment_text
    return text


def _append_reader_warnings(
    state: _ParseState,
    reader_warnings: list[python_warnings.WarningMessage],
) -> None:
    seen: set[str] = set()
    for warning in reader_warnings:
        message = str(warning.message).strip()
        if not message or message in seen:
            continue
        seen.add(message)
        state.quality_warnings.append(
            {
                "type": "xlsx_reader_warning",
                "severity": "low",
                "message": message,
            }
        )


def _parse_worksheet(
    worksheet: Any,
    value_worksheet: Any,
    units: list[EvidenceUnit],
    state: _ParseState,
    *,
    max_cells: int,
    max_tables: int,
) -> None:
    merges, merge_by_position = _worksheet_merges(worksheet, max_cells=max_cells)
    images = list(getattr(worksheet, "_images", []))
    if images:
        state.quality_warnings.append(
            {
                "type": "xlsx_images_unsupported",
                "severity": "medium",
                "message": "Worksheet images are not extracted by the XLSX backend.",
                "sheet_name": worksheet.title,
                "count": len(images),
            }
        )
    charts = list(getattr(worksheet, "_charts", []))
    if charts:
        state.quality_warnings.append(
            {
                "type": "xlsx_charts_unsupported",
                "severity": "medium",
                "message": "Worksheet charts are not extracted by the XLSX backend.",
                "sheet_name": worksheet.title,
                "count": len(charts),
            }
        )

    occupied = _occupied_positions(
        worksheet,
        merges,
        max_cells=max_cells,
    )
    declared_tables, declared_positions = _declared_tables(
        worksheet,
        max_cells=max_cells,
    )
    state.declared_table_count += len(declared_tables)
    if state.declared_table_count > max_tables:
        raise ValueError(
            "XLSX workbook contains too many declared tables: "
            f"{state.declared_table_count} > {max_tables}"
        )
    if not occupied and not declared_tables:
        return

    reader = _SheetReader(worksheet, value_worksheet, state)
    remaining_occupied = occupied - declared_positions
    region_entries = [(table.region, table) for table in declared_tables] + [
        (region, None)
        for region in _logical_regions(
            remaining_occupied,
            barriers=[table.region for table in declared_tables],
        )
    ]
    for region, declared_table in sorted(
        region_entries,
        key=lambda entry: (
            entry[0].min_row,
            entry[0].min_col,
            entry[1] is None,
        ),
    ):
        state.represented_cells += region.area
        if state.represented_cells > max_cells:
            raise ValueError(
                "XLSX workbook contains too many cells in represented ranges: "
                f"{state.represented_cells} > {max_cells}"
            )
        if declared_table is not None:
            if declared_table.header_row_count:
                header_start: int | None = region.min_row
                header_end: int | None = (
                    region.min_row + declared_table.header_row_count - 1
                )
                column_labels: tuple[str, ...] | None = None
            else:
                header_start = None
                header_end = None
                column_labels = declared_table.column_labels
            _emit_table_region(
                worksheet,
                reader,
                region,
                header_start,
                header_end,
                merge_by_position,
                units,
                state,
                column_labels=column_labels,
            )
        else:
            _emit_region(
                worksheet,
                reader,
                region,
                remaining_occupied,
                merge_by_position,
                units,
                state,
            )


def _declared_tables(
    worksheet: Any,
    *,
    max_cells: int,
) -> tuple[list[_DeclaredTable], set[tuple[int, int]]]:
    tables: list[_DeclaredTable] = []
    positions: set[tuple[int, int]] = set()
    represented = 0
    for table in worksheet.tables.values():
        reference = str(getattr(table, "ref", "") or "").strip()
        boundaries = _range_boundaries(reference)
        if boundaries is None:
            raise ValueError(
                f"Invalid Excel table range on sheet {worksheet.title!r}: {reference!r}"
            )
        min_col, min_row, max_col, max_row = boundaries
        if min_row > max_row or min_col > max_col:
            raise ValueError(
                f"Invalid Excel table range on sheet {worksheet.title!r}: {reference!r}"
            )
        region = _Region(min_row, max_row, min_col, max_col)
        if (
            min_row < 1
            or min_col < 1
            or max_row > _EXCEL_MAX_ROWS
            or max_col > _EXCEL_MAX_COLUMNS
        ):
            raise ValueError(
                f"Excel table range is outside worksheet bounds: {reference!r}"
            )
        represented += region.area
        if represented > max_cells:
            raise ValueError(
                "XLSX worksheet contains declared table ranges that exceed "
                "the cell limit"
            )
        for row in range(region.min_row, region.max_row + 1):
            for column in range(region.min_col, region.max_col + 1):
                position = (row, column)
                if position in positions:
                    raise ValueError(
                        f"Overlapping Excel table ranges on sheet {worksheet.title!r}"
                    )
                positions.add(position)

        raw_header_count = getattr(table, "headerRowCount", 1)
        try:
            header_row_count = int(1 if raw_header_count is None else raw_header_count)
        except (TypeError, ValueError) as exc:
            raise ValueError(
                f"Invalid Excel table header count on sheet {worksheet.title!r}"
            ) from exc
        if not 0 <= header_row_count <= region.height:
            raise ValueError(
                f"Invalid Excel table header count on sheet {worksheet.title!r}"
            )

        table_columns = getattr(table, "tableColumns", ()) or ()
        if len(table_columns) > region.width:
            raise ValueError(
                f"Excel table has too many columns on sheet {worksheet.title!r}"
            )
        column_labels = tuple(
            (
                _clean_text(str(getattr(table_columns[offset], "name", "") or ""))
                if offset < len(table_columns)
                else ""
            )
            or _column_letter(region.min_col + offset)
            for offset in range(region.width)
        )
        tables.append(
            _DeclaredTable(
                region=region,
                header_row_count=header_row_count,
                column_labels=column_labels,
            )
        )
    return (
        sorted(tables, key=lambda table: (table.region.min_row, table.region.min_col)),
        positions,
    )


def _regions_overlap(left: _Region, right: _Region) -> bool:
    return not (
        left.max_row < right.min_row
        or right.max_row < left.min_row
        or left.max_col < right.min_col
        or right.max_col < left.min_col
    )


def _worksheet_cells(worksheet: Any) -> Iterator[Any]:
    cells = getattr(worksheet, "_cells", None)
    if isinstance(cells, dict):
        yield from cells.values()
        return
    for row in worksheet.iter_rows():
        yield from row


def _worksheet_merges(
    worksheet: Any,
    *,
    max_cells: int,
) -> tuple[list[_MergedRange], dict[tuple[int, int], _MergedRange]]:
    merges = [
        _MergedRange(
            min_row=int(cell_range.min_row),
            max_row=int(cell_range.max_row),
            min_col=int(cell_range.min_col),
            max_col=int(cell_range.max_col),
        )
        for cell_range in worksheet.merged_cells.ranges
    ]
    merge_by_position: dict[tuple[int, int], _MergedRange] = {}
    represented = 0
    for merged in merges:
        represented += merged.area
        if represented > max_cells:
            raise ValueError(
                "XLSX worksheet contains merged ranges that exceed the cell limit"
            )
        for row in range(merged.min_row, merged.max_row + 1):
            for column in range(merged.min_col, merged.max_col + 1):
                merge_by_position[(row, column)] = merged
    return merges, merge_by_position


def _occupied_positions(
    worksheet: Any,
    merges: list[_MergedRange],
    *,
    max_cells: int,
) -> set[tuple[int, int]]:
    occupied: set[tuple[int, int]] = set()
    for cell in _worksheet_cells(worksheet):
        if not _cell_has_content(cell):
            continue
        occupied.add((int(cell.row), int(cell.column)))
        if len(occupied) > max_cells:
            raise ValueError("XLSX worksheet contains too many occupied cells")
    for merged in merges:
        anchor = worksheet.cell(row=merged.min_row, column=merged.min_col)
        if not _cell_has_content(anchor):
            continue
        for row in range(merged.min_row, merged.max_row + 1):
            for column in range(merged.min_col, merged.max_col + 1):
                occupied.add((row, column))
                if len(occupied) > max_cells:
                    raise ValueError("XLSX worksheet contains too many occupied cells")
    return occupied


def _cell_has_content(cell: Any) -> bool:
    value = getattr(cell, "value", None)
    if value is not None and _clean_text(str(value)):
        return True
    comment = getattr(cell, "comment", None)
    if comment is not None and _clean_text(str(comment.text or "")):
        return True
    hyperlink = getattr(cell, "hyperlink", None)
    if hyperlink is not None:
        return bool(str(hyperlink.target or hyperlink.location or "").strip())
    return False


def _logical_regions(
    occupied: set[tuple[int, int]],
    *,
    barriers: list[_Region] | None = None,
) -> list[_Region]:
    groups = _logical_region_groups(occupied)
    if not barriers:
        return [region for region, _ in groups]

    barrier_index = _RegionSpatialIndex(
        rows=_build_region_interval_index(barriers, axis="row"),
        columns=_build_region_interval_index(barriers, axis="column"),
        row_starts=tuple(sorted(region.min_row for region in barriers)),
        row_ends=tuple(sorted(region.max_row for region in barriers)),
        column_starts=tuple(sorted(region.min_col for region in barriers)),
        column_ends=tuple(sorted(region.max_col for region in barriers)),
    )
    regions: list[_Region] = []
    for region, positions in groups:
        overlapping = _overlapping_indexed_regions(barrier_index, region)
        if not overlapping:
            regions.append(region)
            continue

        row_cuts = {region.min_row, region.max_row + 1}
        for barrier in overlapping:
            row_cuts.add(max(region.min_row, barrier.min_row))
            row_cuts.add(min(region.max_row + 1, barrier.max_row + 1))
        sorted_cuts = sorted(row_cuts)
        positions_by_band: defaultdict[int, set[tuple[int, int]]] = defaultdict(set)
        for position in positions:
            band = bisect_right(sorted_cuts, position[0]) - 1
            positions_by_band[band].add(position)
        for band in sorted(positions_by_band):
            regions.extend(
                fragment
                for fragment, _ in _logical_region_groups(positions_by_band[band])
            )
    return sorted(regions, key=lambda item: (item.min_row, item.min_col))


def _logical_region_groups(
    occupied: set[tuple[int, int]],
) -> list[tuple[_Region, set[tuple[int, int]]]]:
    columns_by_row: defaultdict[int, set[int]] = defaultdict(set)
    for row, column in occupied:
        columns_by_row[row].add(column)
    rows = sorted(columns_by_row)
    groups: list[tuple[_Region, set[tuple[int, int]]]] = []
    for min_row, max_row in _contiguous_ranges(rows):
        band_positions = [
            (row, column)
            for row in range(min_row, max_row + 1)
            for column in columns_by_row[row]
        ]
        band_columns = sorted({column for _, column in band_positions})
        column_ranges = _contiguous_ranges(band_columns)
        range_by_column = {
            column: range_index
            for range_index, (min_col, max_col) in enumerate(column_ranges)
            for column in range(min_col, max_col + 1)
        }
        grouped_positions: defaultdict[int, list[tuple[int, int]]] = defaultdict(list)
        for position in band_positions:
            grouped_positions[range_by_column[position[1]]].append(position)
        for range_index, (min_col, max_col) in enumerate(column_ranges):
            positions = grouped_positions[range_index]
            groups.append(
                (
                    _Region(
                        min_row=min(row for row, _ in positions),
                        max_row=max(row for row, _ in positions),
                        min_col=min(column for _, column in positions),
                        max_col=max(column for _, column in positions),
                    ),
                    set(positions),
                )
            )
    return sorted(
        groups,
        key=lambda item: (item[0].min_row, item[0].min_col),
    )


def _build_region_interval_index(
    regions: list[_Region],
    *,
    axis: str,
) -> _RegionIntervalNode | None:
    if not regions:
        return None
    centers = sorted(sum(_region_axis_bounds(region, axis)) // 2 for region in regions)
    center = centers[len(centers) // 2]
    left: list[_Region] = []
    right: list[_Region] = []
    spanning: list[_Region] = []
    for region in regions:
        start, end = _region_axis_bounds(region, axis)
        if end < center:
            left.append(region)
        elif start > center:
            right.append(region)
        else:
            spanning.append(region)
    by_start = sorted(
        spanning,
        key=lambda region: _region_axis_bounds(region, axis)[0],
    )
    by_end = sorted(
        spanning,
        key=lambda region: _region_axis_bounds(region, axis)[1],
        reverse=True,
    )
    return _RegionIntervalNode(
        center=center,
        by_start=by_start,
        by_end=by_end,
        starts=[_region_axis_bounds(region, axis)[0] for region in by_start],
        negative_ends=[-_region_axis_bounds(region, axis)[1] for region in by_end],
        left=_build_region_interval_index(left, axis=axis),
        right=_build_region_interval_index(right, axis=axis),
    )


def _overlapping_indexed_regions(
    index: _RegionSpatialIndex,
    query: _Region,
) -> list[_Region]:
    row_count = _interval_overlap_count(
        index.row_starts,
        index.row_ends,
        query.min_row,
        query.max_row,
    )
    column_count = _interval_overlap_count(
        index.column_starts,
        index.column_ends,
        query.min_col,
        query.max_col,
    )
    if row_count <= column_count:
        candidates = _query_interval_index(
            index.rows,
            query.min_row,
            query.max_row,
        )
        return [
            region
            for region in candidates
            if not (region.max_col < query.min_col or region.min_col > query.max_col)
        ]
    candidates = _query_interval_index(
        index.columns,
        query.min_col,
        query.max_col,
    )
    return [
        region
        for region in candidates
        if not (region.max_row < query.min_row or region.min_row > query.max_row)
    ]


def _query_interval_index(
    node: _RegionIntervalNode | None,
    query_start: int,
    query_end: int,
) -> list[_Region]:
    if node is None:
        return []
    candidates: list[_Region] = []
    if query_end < node.center:
        stop = bisect_right(node.starts, query_end)
        candidates.extend(node.by_start[:stop])
        candidates.extend(_query_interval_index(node.left, query_start, query_end))
    elif query_start > node.center:
        stop = bisect_right(node.negative_ends, -query_start)
        candidates.extend(node.by_end[:stop])
        candidates.extend(_query_interval_index(node.right, query_start, query_end))
    else:
        candidates.extend(node.by_start)
        candidates.extend(_query_interval_index(node.left, query_start, query_end))
        candidates.extend(_query_interval_index(node.right, query_start, query_end))
    return candidates


def _interval_overlap_count(
    starts: tuple[int, ...],
    ends: tuple[int, ...],
    query_start: int,
    query_end: int,
) -> int:
    ending_before = bisect_right(ends, query_start - 1)
    starting_after = len(starts) - bisect_right(starts, query_end)
    return len(starts) - ending_before - starting_after


def _region_axis_bounds(region: _Region, axis: str) -> tuple[int, int]:
    if axis == "row":
        return region.min_row, region.max_row
    return region.min_col, region.max_col


def _contiguous_ranges(values: list[int]) -> list[tuple[int, int]]:
    if not values:
        return []
    ranges: list[tuple[int, int]] = []
    start = previous = values[0]
    for value in values[1:]:
        if value == previous + 1:
            previous = value
            continue
        ranges.append((start, previous))
        start = previous = value
    ranges.append((start, previous))
    return ranges


def _emit_region(
    worksheet: Any,
    reader: _SheetReader,
    region: _Region,
    occupied: set[tuple[int, int]],
    merge_by_position: dict[tuple[int, int], _MergedRange],
    units: list[EvidenceUnit],
    state: _ParseState,
) -> None:
    if region.width == 1 or region.height == 1:
        _emit_text_region(
            worksheet,
            reader,
            region,
            occupied,
            merge_by_position,
            units,
            state,
        )
        return

    header_bounds = _header_rows(
        worksheet,
        reader,
        region,
        merge_by_position,
    )
    if header_bounds is None:
        _emit_text_region(
            worksheet,
            reader,
            region,
            occupied,
            merge_by_position,
            units,
            state,
        )
        return

    header_start, header_end = header_bounds
    if region.min_row < header_start:
        _emit_text_region(
            worksheet,
            reader,
            _Region(
                min_row=region.min_row,
                max_row=header_start - 1,
                min_col=region.min_col,
                max_col=region.max_col,
            ),
            occupied,
            merge_by_position,
            units,
            state,
        )
    _emit_table_region(
        worksheet,
        reader,
        _Region(
            min_row=header_start,
            max_row=region.max_row,
            min_col=region.min_col,
            max_col=region.max_col,
        ),
        header_start,
        header_end,
        merge_by_position,
        units,
        state,
    )


def _emit_text_region(
    worksheet: Any,
    reader: _SheetReader,
    region: _Region,
    occupied: set[tuple[int, int]],
    merge_by_position: dict[tuple[int, int], _MergedRange],
    units: list[EvidenceUnit],
    state: _ParseState,
) -> None:
    for row in range(region.min_row, region.max_row + 1):
        values: list[str] = []
        anchors: list[tuple[int, int]] = []
        formulas: list[dict[str, str]] = []
        cell_refs: list[dict[str, Any]] = []
        for column in range(region.min_col, region.max_col + 1):
            if (row, column) not in occupied:
                continue
            merged = merge_by_position.get((row, column))
            if merged is not None and merged.anchor != (row, column):
                continue
            payload = reader.cell_payload(
                row,
                column,
                column_id="source",
            )
            source_value = _cell_payload_source_value(payload)
            if source_value:
                values.append(source_value)
            cell = reader.cell(row, column)
            if _cell_has_content(cell):
                anchors.append((row, column))
                cell_ref: dict[str, Any] = {"address": cell.coordinate}
                for key in ("formula", "number_format", "hyperlink", "hidden"):
                    if key in payload:
                        cell_ref[key] = payload[key]
                comments = [
                    str(child.get("content", "")).strip()
                    for child in payload["children"]
                    if str(child.get("content", "")).strip()
                ]
                if comments:
                    cell_ref["comments"] = comments
                cell_refs.append(cell_ref)
            if cell.data_type == "f":
                formulas.append(
                    {
                        "cell": cell.coordinate,
                        "formula": _formula_text(cell.value),
                        "number_format": str(
                            cell.number_format or _GENERAL_NUMBER_FORMAT
                        ),
                    }
                )
        text = " | ".join(values).strip()
        if not text:
            continue

        cell_region = _text_cell_region(
            region,
            anchors,
            merge_by_position,
        )
        cell_range = _range_label(cell_region)
        spreadsheet_metadata: dict[str, Any] = {
            "sheet_name": worksheet.title,
            "sheet_state": worksheet.sheet_state,
            "cell_range": cell_range,
        }
        hidden_rows = [
            source_row
            for source_row in range(cell_region.min_row, cell_region.max_row + 1)
            if worksheet.row_dimensions[source_row].hidden
        ]
        hidden_columns = [
            _column_letter(column)
            for column in range(cell_region.min_col, cell_region.max_col + 1)
            if column in reader.hidden_columns
        ]
        if hidden_rows:
            spreadsheet_metadata["hidden_rows"] = hidden_rows
        if hidden_columns:
            spreadsheet_metadata["hidden_columns"] = hidden_columns
        if cell_refs:
            spreadsheet_metadata["cells"] = cell_refs
        if formulas:
            spreadsheet_metadata["formulas"] = formulas
        units.append(
            EvidenceUnit(
                id=state.next_block_id(),
                type="text",
                format="plain",
                source=SourceEvidence(
                    kind="text",
                    text=_with_sheet_section(worksheet.title, text),
                ),
                content=text,
                metadata={
                    **common_metadata(
                        "text",
                        "plain",
                        section_path=[worksheet.title],
                    ).to_dict(),
                    "spreadsheet": spreadsheet_metadata,
                },
            )
        )


def _text_cell_region(
    region: _Region,
    anchors: list[tuple[int, int]],
    merge_by_position: dict[tuple[int, int], _MergedRange],
) -> _Region:
    if not anchors:
        return region
    min_row = min(row for row, _ in anchors)
    min_col = min(column for _, column in anchors)
    max_row = max(row for row, _ in anchors)
    max_col = max(column for _, column in anchors)
    for anchor in anchors:
        merged = merge_by_position.get(anchor)
        if merged is None:
            continue
        max_row = max(max_row, min(region.max_row, merged.max_row))
        max_col = max(max_col, min(region.max_col, merged.max_col))
    return _Region(min_row, max_row, min_col, max_col)


def _emit_table_region(
    worksheet: Any,
    reader: _SheetReader,
    region: _Region,
    header_start: int | None,
    header_end: int | None,
    merge_by_position: dict[tuple[int, int], _MergedRange],
    units: list[EvidenceUnit],
    state: _ParseState,
    *,
    column_labels: tuple[str, ...] | None = None,
) -> None:
    has_header_rows = header_start is not None and header_end is not None
    columns = [
        {
            "id": f"c{index}",
            "text": (
                _column_header_text(
                    reader,
                    column,
                    header_start,
                    header_end,
                    merge_by_position,
                )
                if has_header_rows
                else (
                    column_labels[index - 1]
                    if column_labels is not None and index <= len(column_labels)
                    else ""
                )
            )
            or _column_letter(column),
        }
        for index, column in enumerate(
            range(region.min_col, region.max_col + 1),
            start=1,
        )
    ]
    header_rows = (
        [
            _structured_row(
                worksheet,
                reader,
                row,
                region,
                merge_by_position,
                span_max_row=header_end,
                row_index=row - header_start + 1,
            )
            for row in range(header_start, header_end + 1)
        ]
        if has_header_rows
        else []
    )
    data_start = header_end + 1 if has_header_rows else region.min_row
    rows = [
        _structured_row(
            worksheet,
            reader,
            row,
            region,
            merge_by_position,
            span_max_row=region.max_row,
            row_index=row - data_start + 1,
        )
        for row in range(data_start, region.max_row + 1)
    ]

    table_id = state.next_table_id()
    headers = [str(column["text"]) for column in columns]
    cell_range = _range_label(region)
    header_range = (
        _range_label(_Region(header_start, header_end, region.min_col, region.max_col))
        if has_header_rows
        else None
    )
    data_range = (
        _range_label(
            _Region(data_start, region.max_row, region.min_col, region.max_col)
        )
        if data_start <= region.max_row
        else None
    )
    content = structured_table(
        caption=worksheet.title,
        columns=columns,
        header_rows=header_rows,
        rows=rows,
    )
    spreadsheet_metadata: dict[str, Any] = {
        "sheet_name": worksheet.title,
        "sheet_state": worksheet.sheet_state,
        "cell_range": cell_range,
        "header_range": header_range,
        "data_range": data_range,
    }
    hidden_rows = [
        source_row
        for source_row in range(region.min_row, region.max_row + 1)
        if worksheet.row_dimensions[source_row].hidden
    ]
    hidden_columns = [
        _column_letter(column)
        for column in range(region.min_col, region.max_col + 1)
        if column in reader.hidden_columns
    ]
    if hidden_rows:
        spreadsheet_metadata["hidden_rows"] = hidden_rows
    if hidden_columns:
        spreadsheet_metadata["hidden_columns"] = hidden_columns

    units.append(
        EvidenceUnit(
            id=state.next_block_id(),
            type="table",
            format="structured_table",
            source=SourceEvidence(
                kind="table",
                text=_table_source_text(
                    worksheet.title,
                    cell_range,
                    columns,
                    header_rows,
                    rows,
                ),
            ),
            content=content,
            metadata={
                **common_metadata(
                    "table",
                    "structured_table",
                    section_path=[worksheet.title],
                ).to_dict(),
                "table": {
                    "table_id": table_id,
                    "headers": headers,
                    "row_count": len(rows),
                },
                "spreadsheet": spreadsheet_metadata,
            },
        )
    )


def _structured_row(
    worksheet: Any,
    reader: _SheetReader,
    source_row: int,
    region: _Region,
    merge_by_position: dict[tuple[int, int], _MergedRange],
    *,
    span_max_row: int,
    row_index: int,
) -> dict[str, Any]:
    cells: list[dict[str, Any]] = []
    for column in range(region.min_col, region.max_col + 1):
        merged = merge_by_position.get((source_row, column))
        if merged is not None and merged.anchor != (source_row, column):
            continue
        rowspan = 1
        colspan = 1
        if merged is not None:
            rowspan = max(
                1,
                min(span_max_row, merged.max_row) - source_row + 1,
            )
            colspan = max(
                1,
                min(region.max_col, merged.max_col) - column + 1,
            )
        column_id = f"c{column - region.min_col + 1}"
        cells.append(
            reader.cell_payload(
                source_row,
                column,
                column_id=column_id,
                rowspan=rowspan,
                colspan=colspan,
            )
        )

    result: dict[str, Any] = {
        "index": row_index,
        "cells": cells,
        "source_row": source_row,
    }
    if worksheet.row_dimensions[source_row].hidden:
        result["hidden"] = True
    return result


def _column_header_text(
    reader: _SheetReader,
    column: int,
    header_start: int,
    header_end: int,
    merge_by_position: dict[tuple[int, int], _MergedRange],
) -> str:
    texts: list[str] = []
    seen_anchors: set[tuple[int, int]] = set()
    for row in range(header_start, header_end + 1):
        merged = merge_by_position.get((row, column))
        anchor = merged.anchor if merged is not None else (row, column)
        if anchor in seen_anchors:
            continue
        seen_anchors.add(anchor)
        text = reader.text(*anchor).strip()
        if text and text not in texts:
            texts.append(text)
    return " / ".join(texts)


def _header_rows(
    worksheet: Any,
    reader: _SheetReader,
    region: _Region,
    merge_by_position: dict[tuple[int, int], _MergedRange],
) -> tuple[int, int] | None:
    filtered = _auto_filter_header_row(worksheet, reader, region, merge_by_position)
    if filtered is not None:
        return filtered, filtered

    search_end = min(region.max_row - 1, region.min_row + 11)
    profiles = {
        row: _row_profile(reader, row, region, merge_by_position)
        for row in range(region.min_row, search_end + 2)
    }
    candidates: list[tuple[int, int]] = []
    for row in range(region.min_row, search_end + 1):
        profile = profiles[row]
        next_profile = profiles[row + 1]
        score = _header_transition_score(profile, next_profile)
        if score:
            candidates.append((score, row))

    fallback_header = _fallback_header_row(profiles, region, search_end)
    if candidates:
        _, transition_header = max(
            candidates,
            key=lambda item: (item[0], -item[1]),
        )
        if fallback_header is None or transition_header == fallback_header:
            header_end = transition_header
        elif transition_header == fallback_header + 1 and _row_has_header_merge(
            fallback_header,
            transition_header,
            region,
            merge_by_position,
        ):
            header_end = transition_header
        elif (
            profiles[transition_header].bold_ratio >= 0.5
            and profiles[fallback_header].bold_ratio < 0.5
        ):
            header_end = transition_header
        else:
            header_end = fallback_header
    else:
        if fallback_header is None:
            return None
        header_end = fallback_header

    header_start = header_end
    while header_start > region.min_row:
        parent_row = header_start - 1
        parent_profile = profiles.get(parent_row)
        has_header_merge = _row_has_header_merge(
            parent_row,
            header_start,
            region,
            merge_by_position,
        )
        if parent_profile is None or not _can_be_parent_header(
            parent_profile,
            allow_single=(
                header_start == header_end
                and _row_has_partial_header_merge(
                    parent_row,
                    header_start,
                    region,
                    merge_by_position,
                )
            ),
        ):
            break
        if not has_header_merge:
            break
        header_start = parent_row
    return header_start, header_end


def _auto_filter_header_row(
    worksheet: Any,
    reader: _SheetReader,
    region: _Region,
    merge_by_position: dict[tuple[int, int], _MergedRange],
) -> int | None:
    auto_filter = getattr(getattr(worksheet, "auto_filter", None), "ref", None)
    if not auto_filter:
        return None
    boundaries = _range_boundaries(str(auto_filter))
    if boundaries is None:
        return None
    min_col, min_row, max_col, max_row = boundaries
    if max_row <= min_row:
        return None
    if not (region.min_row <= min_row <= region.max_row):
        return None
    if max_col < region.min_col or min_col > region.max_col:
        return None
    profile = _row_profile(reader, min_row, region, merge_by_position)
    if profile.total >= 2 and profile.strings >= 2 and profile.formulas == 0:
        return min_row
    return None


def _range_boundaries(reference: str) -> tuple[int, int, int, int] | None:
    try:
        from openpyxl.utils.cell import range_boundaries

        return tuple(int(value) for value in range_boundaries(reference))
    except (TypeError, ValueError):
        return None


def _row_profile(
    reader: _SheetReader,
    row: int,
    region: _Region,
    merge_by_position: dict[tuple[int, int], _MergedRange],
) -> _RowProfile:
    total = strings = non_strings = formulas = bold = 0
    for column in range(region.min_col, region.max_col + 1):
        merged = merge_by_position.get((row, column))
        if merged is not None and merged.anchor != (row, column):
            continue
        cell = reader.cell(row, column)
        if not _cell_has_content(cell):
            continue
        total += 1
        if bool(cell.font.bold):
            bold += 1
        if cell.data_type == "f":
            formulas += 1
            non_strings += 1
        elif _is_text_value(cell.value):
            strings += 1
        else:
            non_strings += 1
    return _RowProfile(
        total=total,
        strings=strings,
        non_strings=non_strings,
        formulas=formulas,
        bold=bold,
    )


def _header_transition_score(
    profile: _RowProfile,
    next_profile: _RowProfile,
) -> int:
    if (
        profile.total < 2
        or profile.strings < 2
        or profile.formulas
        or next_profile.total == 0
    ):
        return 0
    score = 0
    if profile.bold_ratio >= 0.5 and next_profile.bold_ratio < profile.bold_ratio:
        score += 3
    if profile.text_ratio >= 0.75 and (
        next_profile.non_strings > 0 or next_profile.formulas > 0
    ):
        score += 2
    return score


def _fallback_header_row(
    profiles: dict[int, _RowProfile],
    region: _Region,
    search_end: int,
) -> int | None:
    for row in range(region.min_row, search_end + 1):
        profile = profiles[row]
        next_profile = profiles[row + 1]
        if (
            profile.total >= 2
            and profile.strings >= 2
            and profile.formulas == 0
            and next_profile.total > 0
        ):
            return row
    return None


def _can_be_parent_header(
    profile: _RowProfile,
    *,
    allow_single: bool,
) -> bool:
    return (
        profile.total >= (1 if allow_single else 2)
        and profile.strings >= (1 if allow_single else 2)
        and profile.formulas == 0
        and profile.text_ratio >= 0.75
    )


def _row_has_header_merge(
    parent_row: int,
    child_row: int,
    region: _Region,
    merge_by_position: dict[tuple[int, int], _MergedRange],
) -> bool:
    seen: set[tuple[int, int]] = set()
    for column in range(region.min_col, region.max_col + 1):
        merged = merge_by_position.get((parent_row, column))
        if merged is None or merged.anchor in seen:
            continue
        seen.add(merged.anchor)
        if merged.max_col > merged.min_col or merged.max_row >= child_row:
            return True
    return False


def _row_has_partial_header_merge(
    parent_row: int,
    child_row: int,
    region: _Region,
    merge_by_position: dict[tuple[int, int], _MergedRange],
) -> bool:
    seen: set[tuple[int, int]] = set()
    for column in range(region.min_col, region.max_col + 1):
        merged = merge_by_position.get((parent_row, column))
        if merged is None or merged.anchor in seen:
            continue
        seen.add(merged.anchor)
        contributes = merged.max_col > merged.min_col or merged.max_row >= child_row
        covers_full_width = (
            merged.min_col <= region.min_col and merged.max_col >= region.max_col
        )
        if contributes and not covers_full_width:
            return True
    return False


def _table_source_text(
    sheet_name: str,
    cell_range: str,
    columns: list[dict[str, Any]],
    header_rows: list[dict[str, Any]],
    rows: list[dict[str, Any]],
) -> str:
    headers = [str(column["text"]) for column in columns]
    labels = _unique_source_labels(headers)
    lines = [
        f"section: {sheet_name}",
        f"range: {cell_range}",
        f"columns: {' | '.join(labels)}",
    ]
    for header_row in header_rows:
        header_values = [
            f"{cell['address']}={source_value}"
            for cell in header_row["cells"]
            if (source_value := _cell_payload_source_value(cell))
        ]
        if header_values:
            lines.append(
                f"header {header_row['index']} "
                f"[sheet row {header_row['source_row']}]: " + "; ".join(header_values)
            )
    for row in rows:
        values: list[str] = []
        for cell in row["cells"]:
            text = _cell_payload_source_value(cell).strip()
            if not text:
                continue
            column_number = _column_id_number(str(cell["column_id"]))
            label = (
                labels[column_number - 1]
                if 0 < column_number <= len(labels)
                else str(cell["column_id"])
            )
            values.append(f"{label}={text}")
        if values:
            lines.append(
                f"row {row['index']} [sheet row {row['source_row']}]: "
                + "; ".join(values)
            )
    return "\n".join(lines)


def _unique_source_labels(headers: list[str]) -> list[str]:
    totals = Counter(headers)
    seen: dict[str, int] = {}
    labels: list[str] = []
    for header in headers:
        seen[header] = seen.get(header, 0) + 1
        if totals[header] > 1:
            labels.append(f"{header} [{seen[header]}]")
        else:
            labels.append(header)
    return labels


def _column_id_number(column_id: str) -> int:
    match = re.search(r"\d+$", column_id)
    return int(match.group()) if match else 0


def _range_label(region: _Region) -> str:
    return (
        f"{_column_letter(region.min_col)}{region.min_row}:"
        f"{_column_letter(region.max_col)}{region.max_row}"
    )


def _column_letter(column: int) -> str:
    result = ""
    current = column
    while current > 0:
        current, remainder = divmod(current - 1, 26)
        result = chr(ord("A") + remainder) + result
    return result or "A"


def _hidden_column_numbers(worksheet: Any) -> set[int]:
    hidden: set[int] = set()
    for key, dimension in worksheet.column_dimensions.items():
        if not dimension.hidden:
            continue
        min_column = int(dimension.min or _column_number(str(key)))
        max_column = int(dimension.max or min_column)
        if not (1 <= min_column <= max_column <= _EXCEL_MAX_COLUMNS):
            raise ValueError(
                "Worksheet hidden-column range is outside Excel bounds: "
                f"{min_column}:{max_column}"
            )
        hidden.update(range(min_column, max_column + 1))
    return hidden


def _with_sheet_section(sheet_name: str, text: str) -> str:
    return f"section: {sheet_name}\n{text}"


def _formula_text(value: Any) -> str:
    formula_type = value.__class__.__name__
    if formula_type == "ArrayFormula":
        text = _clean_text(str(getattr(value, "text", "") or ""))
        if text:
            return text
        reference = _clean_text(str(getattr(value, "ref", "") or ""))
        return f"=ARRAY_FORMULA(ref={reference})"
    if formula_type == "DataTableFormula":
        parts = [f"ref={_clean_text(str(getattr(value, 'ref', '') or ''))}"]
        for name in ("r1", "r2"):
            field_value = getattr(value, name, None)
            if field_value is not None:
                parts.append(f"{name}={_clean_text(str(field_value))}")
        for name in ("ca", "dt2D", "dtr", "del1", "del2"):
            if bool(getattr(value, name, False)):
                parts.append(f"{name}=TRUE")
        return f"=DATA_TABLE({', '.join(parts)})"
    text = getattr(value, "text", value)
    return _clean_text(str(text or ""))


def _is_text_value(value: Any) -> bool:
    return isinstance(value, str) or value.__class__.__name__ == "CellRichText"


def _format_cell_value(value: Any, number_format: str | None) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.isoformat(timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat(timespec="seconds")
    if isinstance(value, int) and not isinstance(value, bool):
        return _format_number(value, str(number_format or _GENERAL_NUMBER_FORMAT))
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return str(value)
        return _format_number(value, str(number_format or _GENERAL_NUMBER_FORMAT))
    return _clean_text(str(value))


def _format_number(value: int | float, number_format: str) -> str:
    primary_format = number_format.split(";", 1)[0].strip()
    if "%" in primary_format:
        match = re.search(r"0(?:\.(0+))?[^%]*%", primary_format)
        decimals = len(match.group(1) or "") if match else 0
        return f"{value * 100:.{decimals}f}%"
    if re.fullmatch(r"0+", primary_format) and float(value).is_integer():
        return f"{int(value):0{len(primary_format)}d}"
    if isinstance(value, int) or float(value).is_integer():
        return str(int(value))
    return format(float(value), ".15g")


def _clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    return "".join(
        character
        for character in text
        if character in {"\n", "\t"} or ord(character) >= 0x20
    ).strip()
